import pytest

from conftest import make_pdf
from dmx_docs.extract import pymupdf
from dmx_docs.indexer import run_index
from dmx_docs.tools import DocTools

quiet = lambda *a, **k: None  # noqa: E731


@pytest.fixture
def tools(tmp_path, make_cfg):
    root = tmp_path / "RMA"
    (root / "THOR").mkdir(parents=True)
    make_pdf(str(root / "THOR" / "Schema.pdf"), ["Page un : préhenseur.", "Page deux : convoyeur."])
    # A Word file with two embedded pictures (red, then blue).
    from docx import Document
    d = Document()
    d.add_paragraph("Rapport FAT avec photos.")
    for i, color in enumerate(((1, 0, 0), (0, 0, 1))):
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 20), False)
        pix.set_rect(pix.irect, tuple(int(c * 255) for c in color))
        png = tmp_path / f"img{i}.png"
        pix.save(str(png))
        d.add_picture(str(png))
    d.save(str(root / "THOR" / "Rapport FAT.docx"))
    (root / "THOR" / "ancien.doc").write_bytes(b"old word")
    cfg = make_cfg(root)
    run_index(cfg, progress=quiet)
    return DocTools(cfg), root


def test_pdf_page_and_region(tools):
    t, root = tools
    caption, data, fmt = t.view_page(str(root / "THOR" / "Schema.pdf"), page=2)
    assert fmt == "png" and data[:4] == b"\x89PNG"
    assert "page 2/2, whole page" in caption
    w = pymupdf.Pixmap(data).width
    caption, data, _ = t.view_page(str(root / "THOR" / "Schema.pdf"), page=1, region="0,0,0.5,0.5")
    assert "region 0,0,0.5,0.5" in caption and pymupdf.Pixmap(data).width <= w
    with pytest.raises(ValueError, match="between 1 and 2"):
        t.view_page(str(root / "THOR" / "Schema.pdf"), page=3)
    with pytest.raises(ValueError, match="region must be"):
        t.view_page(str(root / "THOR" / "Schema.pdf"), region="top")


def test_docx_pictures_in_order(tools):
    t, root = tools
    for k, rgb in ((1, (255, 0, 0)), (2, (0, 0, 255))):
        caption, data, _ = t.view_page(str(root / "THOR" / "Rapport FAT.docx"), image=k)
        assert f"picture {k} of 2" in caption
        assert tuple(pymupdf.Pixmap(data).pixel(5, 5)) == rgb


def test_works_when_roots_come_from_the_database(tools, make_cfg):
    # Like the MCP server: config.toml has no roots, they are stored by the configuration page.
    t, root = tools
    fresh = DocTools(make_cfg(root, roots=[]))
    fresh.cfg.data_dir = t.cfg.data_dir
    caption, data, _ = fresh.view_page(str(root / "THOR" / "Schema.pdf"))
    assert data[:4] == b"\x89PNG"


def test_refuses_doc_and_paths_outside_the_roots(tools, tmp_path):
    t, root = tools
    with pytest.raises(ValueError, match="cannot be shown"):
        t.view_page(str(root / "THOR" / "ancien.doc"))
    outside = tmp_path / "secret.pdf"
    make_pdf(str(outside), ["x"])
    with pytest.raises(ValueError, match="outside the indexed folders"):
        t.view_page(str(outside))
