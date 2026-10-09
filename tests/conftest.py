import os
import shutil
import subprocess

import pytest

from dmx_docs.config import Config
from dmx_docs.extract import find_libreoffice, pymupdf

SPEC_PAGES = [
    "Spécification de la cellule robotisée Delta\nClient : Nestlé — projet P1234\n"
    "La cellule assure le conditionnement de barres chocolatées en flow-pack.",
    "Système de préhension\nLe préhenseur à ventouses saisit quatre produits par cycle. "
    "Suite à des défauts de prise, la note de modification MN-114 remplace les ventouses "
    "par un préhenseur à doigts souples.",
    "Cadence nominale : 120 produits par minute. Couple de serrage des vis M8 : 25 Nm.",
]


def make_pdf(path, pages):
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=11)
    doc.save(path)
    doc.close()


def make_docx(path):
    from docx import Document

    d = Document()
    d.add_heading("Rapport de mise en service", level=1)
    d.add_paragraph("Le robot a été installé chez le client en mars. Les essais de cadence sont concluants.")
    d.add_heading("Problèmes rencontrés", level=2)
    d.add_paragraph("Vibrations du convoyeur d'alimentation à haute vitesse ; amortisseurs ajoutés.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Pièce"
    table.cell(0, 1).text = "Référence"
    table.cell(1, 0).text = "Amortisseur"
    table.cell(1, 1).text = "AMX-220"
    d.save(path)


def make_png(path, rgb, size=(40, 20)):
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, *size), False)
    pix.set_rect(pix.irect, tuple(rgb))
    pix.save(str(path))


def make_pptx(path, pictures=()):
    """Slide 1: title, subtitle, speaker notes. Slide 2: title, table, text box inside a group.
    Slide 3: the given pictures only (no text)."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])  # Title Slide
    s1.shapes.title.text = "Paloma 4R"
    s1.placeholders[1].text = "Pick and place robot for biscuits"
    s1.notes_slide.notes_text_frame.text = "Mention the washdown version."
    s2 = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    s2.shapes.title.text = "Technical data"
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    for (r, c), text in {(0, 0): "Cadence", (0, 1): "120 ppm", (1, 0): "Robots", (1, 1): "4"}.items():
        table.cell(r, c).text = text
    group = s2.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1)).text_frame.text = "Hygienic design"
    s3 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    for i, picture in enumerate(pictures):
        s3.shapes.add_picture(str(picture), Inches(1 + 3 * i), Inches(1))
    prs.save(str(path))

@pytest.fixture(scope="session")
def corpus(tmp_path_factory):
    root = tmp_path_factory.mktemp("corpus")
    proj = root / "Projets" / "P1234_Nestle"
    proj.mkdir(parents=True)
    make_pdf(str(proj / "Spec_cellule.pdf"), SPEC_PAGES)
    make_docx(str(proj / "Rapport_MES.docx"))
    (root / "Manuals").mkdir()
    make_pdf(str(root / "Manuals" / "maintenance_manual.pdf"),
             ["Maintenance manual. Replace the suction cups every 500 operating hours.",
              "Wartungsanleitung: Die Saugnäpfe alle 500 Betriebsstunden austauschen."])
    (root / "Scans").mkdir()
    make_pdf(str(root / "Scans" / "scan.pdf"), [""])
    (root / "Broken").mkdir()
    (root / "Broken" / "corrupt.pdf").write_bytes(b"%PDF-1.4 this is not really a pdf")
    (root / "~$Rapport_MES.docx").write_bytes(b"lock file")
    (root / "budget.xlsx").write_bytes(b"not indexed")
    soffice = find_libreoffice()
    if soffice:
        old = root / "Old"
        old.mkdir()
        tmp = tmp_path_factory.mktemp("lo")
        make_docx(str(tmp / "ancien_rapport.docx"))
        subprocess.run([soffice, f"-env:UserInstallation=file://{tmp}/profile", "--headless",
                        "--convert-to", "doc", "--outdir", str(old), str(tmp / "ancien_rapport.docx")],
                       check=True, capture_output=True, timeout=180)
    return root


@pytest.fixture
def make_cfg(tmp_path):
    def _make(root, **kw):
        defaults = dict(roots=[str(root)], data_dir=tmp_path / "data", workers=2,
                        embedding_model="hash:256", query_prefix="", passage_prefix="")
        defaults.update(kw)
        return Config(**defaults)
    return _make


@pytest.fixture
def corpus_copy(corpus, tmp_path):
    dst = tmp_path / "corpus"
    shutil.copytree(corpus, dst)
    return dst
