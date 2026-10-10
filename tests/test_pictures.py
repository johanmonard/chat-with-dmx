import hashlib

from conftest import make_png, make_pptx
from dmx_docs.extract import pymupdf
from dmx_docs.pictures import md_context, prepare_document, scale_picture

OPTIONS = {"min_side": 150, "pdf_picture_share": 0.15, "max_pdf_pages": 3000, "doc_converter": "auto",
           "libreoffice_path": None, "excluded_dirs": []}


def job(path, ext, pages=None, native=None, ocr=None, title="Titre", category="Cat"):
    return {"doc_id": 1, "path": str(path), "ext": ext, "title": title, "category": category,
            "pages": pages or {}, "native_chars": native or {}, "ocr_chars": ocr or {}}


def opts(root):
    return dict(OPTIONS, roots=[str(root)])


def by_key(prepared):
    return {(p.page_no, p.image_no): p for p in prepared}


def test_markdown_occurrences_keep_the_marker_numbers(tmp_path):
    (tmp_path / "img").mkdir()
    make_png(tmp_path / "img" / "big.png", (255, 0, 0), size=(300, 200))
    make_png(tmp_path / "img" / "icon.png", (0, 0, 0), size=(40, 20))
    md = tmp_path / "p_FR.md"
    md.write_text("# Manuel\n## Réglage du dépileur\n\nIntroduire l'outil entre les courroies.\n\n"
                  "![🖼 Outil de réglage](img/big.png)\n\n![](missing.png)\n\n![](img/icon.png)\n", encoding="utf-8")
    out = by_key(prepare_document(job(md, ".md"), opts(tmp_path)))
    assert set(out) == {(1, 1), (1, 3)}                      # the broken link 2 is not an occurrence
    first = out[(1, 1)]
    assert first.status == "ready" and first.sha1 == hashlib.sha1((tmp_path / "img" / "big.png").read_bytes()).hexdigest()
    assert "Légende : Outil de réglage" in first.context and "Section : Réglage du dépileur" in first.context
    assert "Introduire l'outil entre les courroies." in first.context and "Document : Titre" in first.context
    assert pymupdf.Pixmap(first.image).width == 300
    assert out[(1, 3)].status == "small" and out[(1, 3)].image is None


def test_markdown_link_outside_the_roots_is_not_an_occurrence(tmp_path):
    root = tmp_path / "root"
    root.mkdir()
    make_png(tmp_path / "secret.png", (0, 255, 0), size=(300, 200))
    md = root / "p.md"
    md.write_text(f"![x]({(tmp_path / 'secret.png').as_posix()})\n", encoding="utf-8")
    assert prepare_document(job(md, ".md"), opts(root)) == []


def test_md_context():
    text = "# T\n## Section A\npara 1\n\n## Section B\n\nLe texte avant.\n\n![a](x.png)\n\n![b](y.png)"
    assert md_context(text, 1) == "Section : Section B\nTexte juste avant l'image : Le texte avant."
    assert md_context(text, 2).startswith("Section : Section B")


def test_word_pictures_in_view_page_order(tmp_path):
    from docx import Document
    from docx.shared import Inches

    make_png(tmp_path / "a.png", (255, 0, 0), size=(300, 200))
    make_png(tmp_path / "b.png", (0, 0, 255), size=(40, 20))
    d = Document()
    d.add_paragraph("Rapport")
    d.add_picture(str(tmp_path / "a.png"), width=Inches(2))
    d.add_picture(str(tmp_path / "b.png"), width=Inches(1))
    d.save(tmp_path / "r.docx")
    out = by_key(prepare_document(job(tmp_path / "r.docx", ".docx"), opts(tmp_path)))
    assert out[(1, 1)].status == "ready" and out[(1, 2)].status == "small"
    assert out[(1, 1)].context == "Document : Titre\nCatégorie : Cat"


def test_powerpoint_pictures_per_slide_with_slide_text(tmp_path):
    make_png(tmp_path / "a.png", (255, 0, 0), size=(300, 200))
    make_png(tmp_path / "b.png", (0, 0, 255), size=(40, 20))
    make_pptx(tmp_path / "d.pptx", pictures=[tmp_path / "a.png", tmp_path / "b.png"])
    out = by_key(prepare_document(job(tmp_path / "d.pptx", ".pptx", pages={3: "Ligne Paloma"}), opts(tmp_path)))
    assert set(out) == {(3, 1), (3, 2)}
    assert out[(3, 1)].status == "ready" and "Texte de la diapositive : Ligne Paloma" in out[(3, 1)].context
    assert out[(3, 2)].status == "small"


def make_picture_pdf(path):
    """p1: text only; p2: text + a picture over half the page; p3: picture only (no text);
    p4: picture only, a scanned letter (OCR found a lot of text)."""
    doc = pymupdf.open()
    photo = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 400, 300), False)
    photo.set_rect(photo.irect, (200, 100, 50))
    for i in range(4):
        page = doc.new_page()
        if i in (0, 1):
            page.insert_textbox(pymupdf.Rect(50, 50, 550, 200), "Brochure Paloma " * 20, fontsize=11)
        if i == 1:
            page.insert_image(pymupdf.Rect(50, 250, 550, 750), pixmap=photo)
        if i in (2, 3):
            page.insert_image(page.rect, pixmap=photo)
    doc.save(path)


def test_pdf_picture_pages(tmp_path):
    make_picture_pdf(tmp_path / "b.pdf")
    j = job(tmp_path / "b.pdf", ".pdf", pages={2: "Brochure Paloma", 3: ""},
            native={1: 400, 2: 400}, ocr={3: 0, 4: 1500})
    out = by_key(prepare_document(j, opts(tmp_path)))
    assert set(out) == {(2, 0), (3, 0)}
    assert out[(2, 0)].status == "ready" and abs(max(scale_picture(out[(2, 0)].image, 5000)[1:]) - 1400) <= 1
    assert "Texte de la page : Brochure Paloma" in out[(2, 0)].context
    no_ocr = by_key(prepare_document(job(tmp_path / "b.pdf", ".pdf", native={1: 400, 2: 400}), opts(tmp_path)))
    assert set(no_ocr) == {(2, 0), (3, 0), (4, 0)}          # OCR off: every page without text is a picture page


def test_picture_file_occurrence(tmp_path):
    (tmp_path / "Pictures" / "Paloma").mkdir(parents=True)
    f = tmp_path / "Pictures" / "Paloma" / "P4R.png"
    make_png(f, (1, 2, 3), size=(300, 200))
    [p] = prepare_document(job(f, ".png", title="P4R", category="Pictures"), opts(tmp_path))
    assert (p.page_no, p.image_no, p.status) == (1, 1, "ready")
    assert p.context == "Fichier : P4R\nDossiers : Pictures / Paloma\nCatégorie : Pictures"


def test_corrupt_picture_file_is_an_error_occurrence(tmp_path):
    f = tmp_path / "bad.jpg"
    f.write_bytes(b"\xff\xd8 not a jpeg")
    [p] = prepare_document(job(f, ".jpg"), opts(tmp_path))
    assert (p.page_no, p.image_no, p.status) == (1, 1, "error") and p.error


def test_unreadable_document_is_one_document_error(tmp_path):
    [p] = prepare_document(job(tmp_path / "gone.pdf", ".pdf"), opts(tmp_path))
    assert (p.page_no, p.image_no, p.status) == (0, 0, "error")


def test_scale_picture_handles_cmyk_and_transparency(tmp_path):
    big = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 3000, 2000), False)
    data, w, h = scale_picture(big.tobytes("png"), 1024)
    assert (w, h) == (3000, 2000) and max(pymupdf.Pixmap(data).width, pymupdf.Pixmap(data).height) == 1024
    cmyk = pymupdf.Pixmap(pymupdf.csCMYK, pymupdf.IRect(0, 0, 200, 200), False)
    assert scale_picture(cmyk.tobytes("jpg"), 1024)[1:] == (200, 200)
    alpha = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 100), True)
    data, _, _ = scale_picture(alpha.tobytes("png"), 1024)
    assert data[:2] == b"\xff\xd8" and not pymupdf.Pixmap(data).alpha       # flattened to a JPEG


def jpeg_with_orientation(width, height, orientation):
    """A JPEG whose EXIF says how to turn it (6 = rotate 90° clockwise to display), like phone photos."""
    import struct

    raw = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, width, height), False).tobytes("jpg")
    ifd = b"II*\x00" + struct.pack("<I", 8) + struct.pack("<H", 1) + struct.pack("<HHIHH", 0x0112, 3, 1, orientation, 0) \
        + struct.pack("<I", 0)
    app1 = b"Exif\x00\x00" + ifd
    return raw[:2] + b"\xff\xe1" + struct.pack(">H", len(app1) + 2) + app1 + raw[2:]


def test_scale_picture_follows_exif_orientation_and_flattens_transparency():
    data, w, h = scale_picture(jpeg_with_orientation(300, 200, 6), 1024)
    shown = pymupdf.Pixmap(data)
    assert (w, h) == (300, 200) and (shown.width, shown.height) == (200, 300)   # turned as a viewer shows it
    clear = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 100), True)
    clear.set_rect(clear.irect, (0, 0, 0, 0))                                 # fully transparent (black + alpha 0)
    data, _, _ = scale_picture(clear.tobytes("png"), 1024)
    flat = pymupdf.Pixmap(data)
    assert not flat.alpha and min(flat.pixel(10, 10)) > 240                  # on white, as view_page shows it


def test_pdf_page_not_yet_read_by_ocr_is_left_for_ocr(tmp_path):
    make_picture_pdf(tmp_path / "b.pdf")
    j = job(tmp_path / "b.pdf", ".pdf", native={1: 400, 2: 400}, ocr={3: 0})   # page 4 not read by OCR yet
    out = by_key(prepare_document(j, dict(opts(tmp_path), ocr_on=True)))
    assert set(out) == {(2, 0), (3, 0)}
