import sys

import pytest

from conftest import make_png, make_pptx
from dmx_docs.extract import extract_file
from dmx_docs.indexer import run_index
from dmx_docs.tools import DocTools

quiet = lambda *a, **k: None  # noqa: E731


def test_pptx_text_per_slide(tmp_path):
    deck = tmp_path / "Deck.pptx"
    make_pptx(deck)
    r = extract_file(str(deck))
    assert r.status == "ok" and r.n_pages == 3
    pages = dict(r.pages)
    assert pages[1].startswith("# Paloma 4R")
    assert "Pick and place robot for biscuits" in pages[1]
    assert "Notes: Mention the washdown version." in pages[1]
    assert "Cadence | 120 ppm" in pages[2] and "Robots | 4" in pages[2]
    assert "Hygienic design" in pages[2]           # text box inside a group
    assert pages[2].count("Technical data") == 1   # the title is not repeated
    assert 3 not in pages                          # pictures only: no text


def test_broken_pptx_is_an_error_not_a_crash(tmp_path):
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"not a zip file")  # also what an encrypted (password) .pptx looks like to python-pptx
    r = extract_file(str(bad))
    assert r.status == "error" and r.error


@pytest.fixture
def deck_tools(tmp_path, make_cfg):
    root = tmp_path / "Marketing"
    (root / "Presentations").mkdir(parents=True)
    red, blue = tmp_path / "red.png", tmp_path / "blue.png"
    make_png(red, (255, 0, 0))
    make_png(blue, (0, 0, 255))
    make_pptx(root / "Presentations" / "Deck.pptx", pictures=(red, blue))
    cfg = make_cfg(root, extensions=[".pdf", ".pptx", ".ppt"], world="marketing", profile="marketing")
    run_index(cfg, progress=quiet)
    return DocTools(cfg), root / "Presentations" / "Deck.pptx"


def test_pptx_is_searched_and_read_by_slide(deck_tools):
    t, deck = deck_tools
    assert "Deck.pptx — slide 2/3 (PowerPoint" in t.search("Hygienic", mode="keyword")
    text = t.read_document(str(deck))
    assert "3 slides" in text and "--- slide 2 ---" in text and "[no text on this page]" in text


def test_pptx_table_keeps_repeated_values_and_merged_cells(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    deck = tmp_path / "Table.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    table = slide.shapes.add_table(2, 3, Inches(1), Inches(1), Inches(6), Inches(2)).table
    table.cell(0, 0).merge(table.cell(0, 1))  # header: the first two cells are merged
    table.cell(0, 0).text = "A B"
    table.cell(0, 2).text = "C"
    for c, text in enumerate(["Robots", "4", "4"]):  # equal adjacent values are data, not repeats
        table.cell(1, c).text = text
    prs.save(str(deck))
    r = extract_file(str(deck))
    assert r.status == "ok"
    text = dict(r.pages)[1]
    assert "A B | C" in text.splitlines()
    assert "Robots | 4 | 4" in text.splitlines()
    assert "|  |" not in text  # no empty cell where the merged cell was


def test_pptx_table_keeps_columns_under_a_vertical_merge(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    deck = tmp_path / "Vertical.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(6), Inches(3)).table
    table.cell(0, 0).merge(table.cell(1, 0))  # "Group" spans the first two rows of the first column
    table.cell(0, 0).text = "Group"
    table.cell(0, 1).text = "Speed"
    table.cell(1, 1).text = "Max"
    table.cell(2, 0).text = "Motor"
    table.cell(2, 1).text = "3 kW"
    prs.save(str(deck))
    r = extract_file(str(deck))
    assert r.status == "ok"
    lines = dict(r.pages)[1].splitlines()
    assert "Group | Speed" in lines
    assert "| Max" in lines  # Max stays in the second column: the covered cell leaves an empty slot
    assert "Motor | 3 kW" in lines
